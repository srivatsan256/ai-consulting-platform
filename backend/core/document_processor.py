import os
from PyPDF2 import PdfReader # type: ignore
from docx import Document as DocxDocument
from openpyxl import load_workbook  # type: ignore


def extract_text_from_pdf(file_path: str) -> dict:
    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return {
        "text": "\n\n".join(text_parts),
        "total_pages": len(reader.pages),
    }


def extract_text_from_docx(file_path: str) -> dict:
    doc = DocxDocument(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    return {
        "text": "\n\n".join(text_parts),
        "total_pages": 0,
    }


def extract_text_from_xlsx(file_path: str) -> dict:
    workbook = load_workbook(file_path, read_only=True, data_only=True)
    text_parts = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                text_parts.append(" | ".join(cells))
    workbook.close()
    return {
        "text": "\n\n".join(text_parts),
        "total_pages": 0,
    }


def extract_text_from_file(file_path: str) -> dict:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext in (".docx", ".doc"):
        return extract_text_from_docx(file_path)
    elif ext in (".xlsx", ".xls"):
        return extract_text_from_xlsx(file_path)
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8") as f:
            return {"text": f.read(), "total_pages": 0}
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end]
        chunks.append(chunk)
        start = end - overlap
    return chunks


def extract_text_from_uploaded_file(uploaded_file) -> dict:
    import tempfile
    suffix = ""
    if hasattr(uploaded_file, "name"):
        suffix = os.path.splitext(uploaded_file.name)[1]
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        for chunk in uploaded_file.chunks():
            tmp.write(chunk)
        tmp_path = tmp.name
    try:
        result = extract_text_from_file(tmp_path)
    finally:
        os.unlink(tmp_path)
    return result
